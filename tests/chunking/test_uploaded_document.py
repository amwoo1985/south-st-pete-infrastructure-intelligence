"""Tests for app/chunking/uploaded_document.py — the POST
/documents/upload chunker (Phase H). Exercises the real
chunk_uploaded_document() entry point (not private helpers standing in
for it), against real TXT/DOCX bytes built in-test and a real PDF
fixture already present in tests/fixtures/legistar/, matching this
repo's real-input-over-hand-built-stub testing convention.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.chunking.uploaded_document import (
    DOC_TYPE,
    MAX_SAFE_CHUNK_TEXT_CHARS,
    EmptyDocumentError,
    chunk_uploaded_document,
)

FIXTURES_ROOT = Path(__file__).parent.parent / "fixtures"
FILE_HASH = "deadbeef" * 8  # fixed, arbitrary 64-hex-char stand-in for a real sha256


def _txt_bytes(*paragraphs: str) -> bytes:
    return "\n\n".join(paragraphs).encode("utf-8")


def _docx_bytes(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """paragraphs: list of (text, style_name_or_None) — style_name e.g.
    "Heading 1" to produce a heading-styled paragraph, matching what a
    real DOCX's built-in styles produce."""
    document = DocxDocument()
    for text, style in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


# --- Size-cap grouping (TXT path exercises the plain fallback grouping) -----


def test_paragraphs_under_cap_group_into_one_chunk():
    raw = _txt_bytes("Paragraph one.", "Paragraph two.", "Paragraph three.")
    chunks = chunk_uploaded_document(
        raw, filename="doc.txt", file_hash=FILE_HASH, kind="txt", target_chars=2000
    )
    assert len(chunks) == 1
    assert chunks[0].text == "Paragraph one.\n\nParagraph two.\n\nParagraph three."
    assert chunks[0].doc_type == DOC_TYPE
    assert chunks[0].section_label == "doc.txt"
    assert chunks[0].attribution.source_url == f"upload://{FILE_HASH}"
    assert chunks[0].attribution.published_date is None


def test_paragraphs_split_into_multiple_groups_once_cap_exceeded():
    # Each paragraph is 100 chars; cap of 250 fits 2 per group (100 + 2
    # (join) + 100 = 202 <= 250; a 3rd would be 304 > 250).
    para = "x" * 100
    raw = _txt_bytes(para, para, para, para, para)
    chunks = chunk_uploaded_document(
        raw, filename="doc.txt", file_hash=FILE_HASH, kind="txt", target_chars=250
    )
    assert len(chunks) == 3  # [p,p] [p,p] [p]
    assert [c.text.count("\n\n") + 1 for c in chunks] == [2, 2, 1]
    # No paragraph content lost or duplicated across groups.
    assert sum(c.text.count("x") for c in chunks) == 500
    # Multiple chunks from the same (headingless) section -> filename +
    # part index, never a bare duplicate label.
    assert [c.section_label for c in chunks] == [
        "doc.txt (part 1)",
        "doc.txt (part 2)",
        "doc.txt (part 3)",
    ]


def test_oversized_single_paragraph_is_never_split():
    oversized = "y" * 3000  # alone, already over a 2000-char cap
    raw = _txt_bytes(oversized)
    chunks = chunk_uploaded_document(
        raw, filename="doc.txt", file_hash=FILE_HASH, kind="txt", target_chars=2000
    )
    assert len(chunks) == 1
    assert chunks[0].text == oversized  # intact, not truncated or split


def test_oversized_paragraph_does_not_absorb_the_next_paragraph():
    oversized = "y" * 3000
    small = "small paragraph"
    raw = _txt_bytes(oversized, small)
    chunks = chunk_uploaded_document(
        raw, filename="doc.txt", file_hash=FILE_HASH, kind="txt", target_chars=2000
    )
    assert len(chunks) == 2
    assert chunks[0].text == oversized
    assert chunks[1].text == small


# --- DOCX heading-based section boundaries -----------------------------------


def test_docx_headings_create_section_boundaries():
    raw = _docx_bytes(
        [
            ("Intro", "Heading 1"),
            ("Intro paragraph text.", None),
            ("Details", "Heading 1"),
            ("Details paragraph text.", None),
        ]
    )
    chunks = chunk_uploaded_document(
        raw, filename="doc.docx", file_hash=FILE_HASH, kind="docx", target_chars=2000
    )
    assert len(chunks) == 2
    assert chunks[0].section_label == "Intro"
    # Heading text is folded into the embedded text itself, not just the
    # label (see module docstring) — same as
    # app/chunking/stpete_pages.py's h2-per-chunk precedent.
    assert chunks[0].text == "Intro\n\nIntro paragraph text."
    assert chunks[1].section_label == "Details"
    assert chunks[1].text == "Details\n\nDetails paragraph text."
    assert all(c.doc_type == DOC_TYPE for c in chunks)


def test_docx_content_before_first_heading_is_its_own_section():
    raw = _docx_bytes(
        [
            ("Preamble text before any heading.", None),
            ("Details", "Heading 1"),
            ("Details paragraph text.", None),
        ]
    )
    chunks = chunk_uploaded_document(
        raw, filename="doc.docx", file_hash=FILE_HASH, kind="docx", target_chars=2000
    )
    assert len(chunks) == 2
    assert chunks[0].section_label == "doc.docx"  # no heading -> filename fallback
    assert chunks[0].text == "Preamble text before any heading."
    assert chunks[1].section_label == "Details"


def test_docx_without_any_headings_falls_back_to_filename_label():
    raw = _docx_bytes([("Just a plain paragraph, no heading styles at all.", None)])
    chunks = chunk_uploaded_document(
        raw, filename="plain.docx", file_hash=FILE_HASH, kind="docx", target_chars=2000
    )
    assert len(chunks) == 1
    assert chunks[0].section_label == "plain.docx"


def test_docx_oversized_heading_section_still_splits_by_size_not_split_paragraphs():
    para = "z" * 100
    raw = _docx_bytes(
        [("Big Section", "Heading 1")] + [(para, None)] * 5
    )
    chunks = chunk_uploaded_document(
        raw, filename="doc.docx", file_hash=FILE_HASH, kind="docx", target_chars=250
    )
    # Same 2-per-group math as the TXT cap test above, but now within one
    # heading's section -> "(part N)" suffix on the heading label itself.
    assert len(chunks) == 3
    assert [c.section_label for c in chunks] == [
        "Big Section (part 1)",
        "Big Section (part 2)",
        "Big Section (part 3)",
    ]


# --- Fail-loud on empty documents (.claude/rules/crawler.md) ----------------


def test_whitespace_only_txt_raises_empty_document_error():
    raw = "   \n\n\t\n  ".encode("utf-8")
    with pytest.raises(EmptyDocumentError):
        chunk_uploaded_document(raw, filename="blank.txt", file_hash=FILE_HASH, kind="txt")


def test_docx_with_only_blank_paragraphs_raises_empty_document_error():
    raw = _docx_bytes([("", None), ("   ", None)])
    with pytest.raises(EmptyDocumentError):
        chunk_uploaded_document(raw, filename="blank.docx", file_hash=FILE_HASH, kind="docx")


# --- Deterministic / idempotent chunk_ids ------------------------------------


def test_chunk_ids_are_deterministic_across_repeated_calls():
    raw = _txt_bytes("Paragraph one.", "Paragraph two.")
    first = chunk_uploaded_document(raw, filename="doc.txt", file_hash=FILE_HASH, kind="txt")
    second = chunk_uploaded_document(raw, filename="doc.txt", file_hash=FILE_HASH, kind="txt")
    assert [c.chunk_id for c in first] == [c.chunk_id for c in second]
    # Never derived from a random/incrementing source (app/chunking/base.py
    # contract) — a different file_hash must change the id even for
    # byte-identical content.
    other_hash_chunks = chunk_uploaded_document(
        raw, filename="doc.txt", file_hash="a" * 64, kind="txt"
    )
    assert [c.chunk_id for c in first] != [c.chunk_id for c in other_hash_chunks]


def test_unsupported_kind_raises_value_error():
    with pytest.raises(ValueError):
        chunk_uploaded_document(b"content", filename="doc.xyz", file_hash=FILE_HASH, kind="xyz")


# --- Real PDF fixture smoke test ---------------------------------------------


def test_real_pdf_fixture_produces_nonempty_chunks():
    raw = (FIXTURES_ROOT / "legistar" / "agenda_fallback.pdf").read_bytes()
    chunks = chunk_uploaded_document(
        raw, filename="agenda_fallback.pdf", file_hash=FILE_HASH, kind="pdf"
    )
    assert len(chunks) >= 1
    assert all(c.doc_type == DOC_TYPE for c in chunks)
    assert all(c.text.strip() for c in chunks)


# --- PDF page-boundary paragraph merging (rag-review pre-commit fix) --------


def test_pdf_paragraph_spanning_a_page_break_is_not_split(monkeypatch):
    """Regression test for the rag-review pre-commit finding: PDF
    paragraph extraction used to split per-page BEFORE flattening, which
    force-split any clause spanning a page break (no blank line at the
    seam) into two separate "paragraphs" before _group_paragraphs_by_size
    ever saw them — defeating the never-split-mid-thought guarantee at
    exactly the boundary it exists to protect. _paragraphs_from_pdf now
    joins all pages' text into one string first and splits once,
    globally."""

    class _FakePage:
        def __init__(self, text: str) -> None:
            self._text = text

        def extract_text(self) -> str:
            return self._text

    class _FakeReader:
        def __init__(self, _stream) -> None:
            # A real-shaped case: a sentence runs right up to the end of
            # page one and continues on page two with no blank line at
            # the seam; a second, genuinely separate paragraph follows.
            self.pages = [
                _FakePage("This sentence starts on page one and continues"),
                _FakePage(
                    "onto page two with no blank line at the seam.\n\n"
                    "A second, separate paragraph."
                ),
            ]

    monkeypatch.setattr("app.chunking.uploaded_document.PdfReader", _FakeReader)

    # target_chars=10 forces each real paragraph into its own group (both
    # are already over such a tiny cap) so each chunk's text maps 1:1 to
    # one real paragraph, making the page-break-merge behavior directly
    # observable.
    chunks = chunk_uploaded_document(
        b"fake-pdf-bytes", filename="split.pdf", file_hash=FILE_HASH, kind="pdf", target_chars=10
    )
    assert len(chunks) == 2
    assert chunks[0].text == (
        "This sentence starts on page one and continues\n"
        "onto page two with no blank line at the seam."
    )
    assert chunks[1].text == "A second, separate paragraph."


# --- Oversized-paragraph safety-ceiling fallback (rag-review pre-commit fix) -


def test_oversized_paragraph_beyond_safety_ceiling_gets_sentence_split():
    sentence = "This is one real sentence with real words in it. "
    # Well beyond MAX_SAFE_CHUNK_TEXT_CHARS, built from many real
    # sentences so sentence-boundary splitting has plenty of usable
    # boundaries — the preferred path, not the hard-character fallback.
    repeat_count = MAX_SAFE_CHUNK_TEXT_CHARS // len(sentence) + 50
    huge_paragraph = sentence * repeat_count
    assert len(huge_paragraph) > MAX_SAFE_CHUNK_TEXT_CHARS

    raw = _txt_bytes(huge_paragraph)
    chunks = chunk_uploaded_document(
        raw, filename="huge.txt", file_hash=FILE_HASH, kind="txt", target_chars=2000
    )

    assert len(chunks) > 1
    assert all(len(c.text) <= MAX_SAFE_CHUNK_TEXT_CHARS for c in chunks)
    # Whole-sentence pieces, not an arbitrary mid-sentence cut.
    assert all(c.text.strip().endswith(".") for c in chunks)
    # Labeled distinctly from a normal size-grouping "(part N)" split, so
    # a citation makes clear this was a degraded fallback split.
    assert [c.section_label for c in chunks] == [
        f"huge.txt (split {i}/{len(chunks)})" for i in range(1, len(chunks) + 1)
    ]
    assert len(set(c.chunk_id for c in chunks)) == len(chunks)


def test_oversized_paragraph_with_no_sentence_boundaries_hard_splits():
    # No '.', '!', or '?' anywhere — no usable sentence boundary at all,
    # forcing the true last-resort fixed-size character split.
    huge_paragraph = "w" * (MAX_SAFE_CHUNK_TEXT_CHARS + 1000)
    raw = _txt_bytes(huge_paragraph)
    chunks = chunk_uploaded_document(
        raw, filename="wall.txt", file_hash=FILE_HASH, kind="txt", target_chars=2000
    )

    assert len(chunks) == 2
    assert all(len(c.text) <= MAX_SAFE_CHUNK_TEXT_CHARS for c in chunks)
    # Lossless partition — every character preserved, none duplicated.
    assert chunks[0].text + chunks[1].text == huge_paragraph
    assert len(set(c.chunk_id for c in chunks)) == len(chunks)


def test_paragraph_under_safety_ceiling_is_never_run_through_the_fallback():
    # Sanity check on the other side of the boundary: a paragraph that's
    # oversized relative to CHUNK_TARGET_SIZE_CHARS but still comfortably
    # under MAX_SAFE_CHUNK_TEXT_CHARS must come back as exactly one
    # chunk, with the plain (non-"split") label.
    paragraph = "y" * (MAX_SAFE_CHUNK_TEXT_CHARS - 100)
    raw = _txt_bytes(paragraph)
    chunks = chunk_uploaded_document(
        raw, filename="fits.txt", file_hash=FILE_HASH, kind="txt", target_chars=2000
    )
    assert len(chunks) == 1
    assert chunks[0].text == paragraph
    assert chunks[0].section_label == "fits.txt"


# --- DOCX heading-detection silent-degradation visibility (rag-review follow-up) -


def test_docx_without_headings_logs_debug_for_a_substantial_document(caplog):
    paragraphs = [
        (f"Paragraph number {i} with some real content in it.", None) for i in range(8)
    ]
    raw = _docx_bytes(paragraphs)
    with caplog.at_level("DEBUG", logger="chunking.uploaded_document"):
        chunk_uploaded_document(
            raw, filename="no_headings.docx", file_hash=FILE_HASH, kind="docx"
        )
    assert any(
        "zero `Heading N`-styled paragraphs detected" in record.message
        for record in caplog.records
    )
